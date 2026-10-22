"""In-memory screening for employer-supplied expense exports.

The analyser reports data-quality issues and explainable transaction signals.
It never claims to verify a payment, receipt, employee identity or fraud.
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile
from pathlib import PurePath
from statistics import median
from typing import Any

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MIN_PEER_COUNT = 8
ROBUST_Z_THRESHOLD = 3.5

ALIASES = {
    "transaction_id": {"transactionid", "transactionreference", "transactionref", "id", "reference", "ref", "recordid"},
    "transaction_date": {"transactiondate", "date", "posteddate", "purchasedate", "bookingdate", "valuedate"},
    "employee_id": {"employeeid", "employee", "staffid", "cardholderid", "user", "owner", "person"},
    "department": {"department", "costcenter", "costcentre", "businessunit", "team"},
    "merchant_name": {"merchant", "merchantname", "vendor", "vendorname", "supplier", "suppliername", "payee", "description", "counterparty"},
    "category": {"category", "expensecategory", "spendcategory", "transactiontype", "type", "mccdescription"},
    "amount": {"amount", "transactionamount", "expenseamount", "total", "value", "debit", "grossamount"},
    "currency": {"currency", "currencycode", "ccy"},
    "payment_method": {"paymentmethod", "method", "cardtype", "instrument"},
    "business_purpose": {"businesspurpose", "purpose", "justification", "reason"},
    "receipt_reference": {"receiptreference", "receiptid", "invoicereference", "invoicenumber", "receipt"},
}

OPTIONAL_FIELDS = ("employee_id", "department", "category", "currency", "business_purpose", "receipt_reference", "payment_method")
REQUIRED_FIELDS = ("transaction_date", "merchant_name", "amount")


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _parse_amount(value: Any) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError("amount is blank")
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        raw = str(value).strip().replace(" ", "")
        negative = raw.startswith("(") and raw.endswith(")")
        raw = raw.strip("()")
        raw = re.sub(r"[^0-9,.\-]", "", raw)
        if "," in raw and "." in raw:
            if raw.rfind(",") > raw.rfind("."):
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", "")
        elif "," in raw:
            tail = raw.rsplit(",", 1)[-1]
            raw = raw.replace(",", ".") if len(tail) in (1, 2) else raw.replace(",", "")
        if negative and not raw.startswith("-"):
            raw = "-" + raw
        amount = Decimal(raw)
    if not amount.is_finite():
        raise ValueError("amount is not a finite number")
    return amount.quantize(Decimal("0.01"))


def _parse_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = str(value or "").strip()
    if not raw:
        raise ValueError("date is blank")
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError("date format is not recognised")


def _read_rows(filename: str, raw: bytes) -> tuple[list[str], list[tuple[int, dict]]]:
    extension = PurePath(filename).suffix.lower()
    if extension == ".csv":
        text = raw.decode("utf-8-sig")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("The CSV file needs a header row.")
        headers = [str(x or "").strip() for x in reader.fieldnames]
        if len(headers) > 100:
            raise ValueError("The file has more than 100 columns. Remove unused columns and retry.")
        rows = [(line_no, dict(row)) for line_no, row in enumerate(reader, start=2)
                if any(str(v or "").strip() for v in row.values())]
        if len(rows) > MAX_ROWS:
            raise ValueError(f"File contains more than {MAX_ROWS:,} data rows. Split it into smaller files.")
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX support is unavailable. Install the openpyxl dependency.") from exc
        try:
            with ZipFile(io.BytesIO(raw)) as archive:
                expanded_size = sum(item.file_size for item in archive.infolist())
                if expanded_size > 50 * 1024 * 1024:
                    raise ValueError("The Excel workbook expands beyond the 50 MB safety limit.")
        except BadZipFile as exc:
            raise ValueError("The XLSX file is not a valid Excel workbook.") from exc
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header_values = next(iterator, None)
            if not header_values:
                raise ValueError("The Excel worksheet needs a header row.")
            headers = [str(x or "").strip() for x in header_values]
            if len(headers) > 100:
                raise ValueError("The worksheet has more than 100 columns. Remove unused columns and retry.")
            rows = []
            for line_no, values in enumerate(iterator, start=2):
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                rows.append((line_no, {headers[i]: values[i] if i < len(headers) else None
                                       for i in range(len(headers))}))
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"File contains more than {MAX_ROWS:,} data rows. Split it into smaller files.")
        finally:
            workbook.close()
    else:
        raise ValueError("Unsupported file type. Upload a .csv or .xlsx file.")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"File contains more than {MAX_ROWS:,} data rows. Split it into smaller files.")
    return headers, rows


def _map_columns(headers: list[str]) -> dict[str, str | None]:
    by_key = {_header_key(header): header for header in headers if _header_key(header)}
    mapping: dict[str, str | None] = {}
    for field, aliases in ALIASES.items():
        mapping[field] = next((by_key[a] for a in aliases if a in by_key), None)
    missing = [field for field in REQUIRED_FIELDS if mapping.get(field) is None]
    if missing:
        descriptions = {
            "transaction_date": "a date column (Date, Transaction Date, Posted Date)",
            "merchant_name": "a merchant/payee/description column",
            "amount": "an amount/value column",
        }
        raise ValueError("Could not identify " + "; ".join(descriptions[x] for x in missing) +
                         ". Rename the source columns and upload again.")
    return mapping


def _normalise_rows(raw_rows: list[tuple[int, dict]], mapping: dict[str, str | None]) -> tuple[list[dict], list[dict]]:
    accepted: list[dict] = []
    errors: list[dict] = []
    for line_no, source in raw_rows:
        def get(field: str) -> Any:
            column = mapping.get(field)
            return source.get(column) if column else None
        try:
            tx_date = _parse_date(get("transaction_date"))
            amount = _parse_amount(get("amount"))
            if amount == 0:
                raise ValueError("amount is zero")
            merchant = str(get("merchant_name") or "").strip()
            if not merchant:
                raise ValueError("merchant/payee is blank")
            currency = str(get("currency") or "UNKNOWN").strip().upper()
            if currency != "UNKNOWN" and not re.fullmatch(r"[A-Z]{3}", currency):
                raise ValueError("currency must be a three-letter code")
            category = str(get("category") or "Uncategorised").strip()
            employee = str(get("employee_id") or "").strip()
            department = str(get("department") or "").strip()
            source_id = str(get("transaction_id") or "").strip()
            stable = json_safe([tx_date, merchant.casefold(), str(amount), currency, employee, source_id, line_no])
            transaction_id = source_id or "UP-" + hashlib.sha256(stable.encode()).hexdigest()[:14].upper()
            accepted.append({
                "transaction_id": transaction_id,
                "source_transaction_id": source_id or None,
                "transaction_date": tx_date,
                "employee_id": employee or None,
                "department": department or None,
                "merchant_name": merchant,
                "merchant_key": re.sub(r"\s+", " ", merchant.casefold()).strip(),
                "category": category,
                "amount": float(amount),
                "currency": currency,
                "payment_method": str(get("payment_method") or "").strip() or None,
                "business_purpose": str(get("business_purpose") or "").strip() or None,
                "receipt_reference": str(get("receipt_reference") or "").strip() or None,
                "source_row": line_no,
                "signals": [],
                "review_status": "NO_FLAG",
            })
            if tx_date > date.today().isoformat():
                accepted[-1]["signals"].append({
                    "type": "FUTURE_DATE", "severity": "medium",
                    "title": "Transaction date is in the future",
                    "explanation": "The recorded transaction date is later than today. Check whether the source uses a planned date or an incorrect date format.",
                    "evidence": {"transaction_date": tx_date},
                })
                accepted[-1]["review_status"] = "REVIEW"
        except (ValueError, InvalidOperation) as exc:
            errors.append({"row": line_no, "reason": str(exc)})
    return accepted, errors


def json_safe(value: Any) -> str:
    return repr(value)


def analyse_rows(rows: list[dict], filename: str, mapping: dict[str, str | None],
                 row_errors: list[dict]) -> dict:
    duplicates: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        key = (
            (row["employee_id"] or "").casefold(),
            row["merchant_key"],
            row["transaction_date"],
            round(row["amount"], 2),
            row["currency"],
        )
        duplicates[key].append(row)

    duplicate_groups = 0
    for members in duplicates.values():
        if len(members) > 1:
            duplicate_groups += 1
            for row in members:
                row["signals"].append({
                    "type": "POSSIBLE_DUPLICATE",
                    "severity": "high",
                    "title": "Possible duplicate transaction",
                    "explanation": "The same merchant, date, amount and currency appear more than once. Confirm the source transaction references before treating these as duplicates.",
                    "evidence": {"matching_rows": [m["source_row"] for m in members]},
                })

    cohorts: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        if row["amount"] > 0:
            cohorts[(row["merchant_key"], row["category"].casefold(), row["currency"])].append(row)

    outlier_count = 0
    insufficient_cohorts = 0
    for members in cohorts.values():
        if len(members) - 1 < MIN_PEER_COUNT:
            insufficient_cohorts += 1
            continue
        amounts = [m["amount"] for m in members]
        centre = median(amounts)
        deviations = [abs(v - centre) for v in amounts]
        mad = median(deviations)
        scale = 1.4826 * mad if mad else max(abs(centre) * 0.05, 1.0)
        for row in members:
            score = (row["amount"] - centre) / scale
            if score >= ROBUST_Z_THRESHOLD:
                outlier_count += 1
                row["signals"].append({
                    "type": "UNUSUAL_AMOUNT",
                    "severity": "medium",
                    "title": "Unusual amount within merchant/category history",
                    "explanation": f"Amount {row['amount']:.2f} is materially above this file's comparison median of {centre:.2f} (robust score {score:.2f}). This is a screening signal, not evidence of wrongdoing.",
                    "evidence": {"amount": round(row["amount"], 2), "cohort_median": round(centre, 2),
                                 "comparison_rows": len(members), "robust_score": round(score, 2)},
                })

    missing_optional = [field for field in OPTIONAL_FIELDS if not mapping.get(field)]
    for row in rows:
        if row["amount"] < 0:
            row["signals"].append({
                "type": "CREDIT_OR_REFUND",
                "severity": "info",
                "title": "Credit or refund entry",
                "explanation": "Negative values are retained separately from positive expense totals. Confirm whether this is a refund, reversal or source-format convention.",
                "evidence": {"amount": row["amount"]},
            })
        if row["signals"]:
            row["review_status"] = "REVIEW"
        row["signal_count"] = len(row["signals"])

    currencies: dict[str, float] = defaultdict(float)
    for row in rows:
        if row["amount"] > 0:
            currencies[row["currency"]] += row["amount"]

    signal_rows = sum(1 for row in rows if row["review_status"] == "REVIEW")
    return {
        "filename": PurePath(filename).name,
        "analysis_type": "transaction screening",
        "summary": {
            "rows_received": len(rows) + len(row_errors),
            "rows_analysed": len(rows),
            "rows_with_signals": signal_rows,
            "possible_duplicate_groups": duplicate_groups,
            "unusual_amount_signals": outlier_count,
            "invalid_rows": len(row_errors),
            "positive_expense_totals": {k: round(v, 2) for k, v in sorted(currencies.items())},
        },
        "data_coverage": {
            "columns_detected": {k: v for k, v in mapping.items() if v},
            "optional_fields_not_supplied": missing_optional,
            "peer_groups_without_enough_history": insufficient_cohorts,
        },
        "limitations": [
            "A transaction export alone cannot prove that a payment occurred or that a receipt is genuine.",
            "No employer-specific policy was supplied with this file, so policy compliance is not being certified.",
            "A flag is a review lead, not a fraud finding. An unflagged row is not automatically approved.",
        ],
        "row_errors": row_errors[:100],
        "items": rows,
    }


def analyse_upload(filename: str, raw: bytes) -> dict:
    if not filename:
        raise ValueError("Choose a file to upload.")
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("The file is larger than 10 MB. Split it into smaller files.")
    headers, raw_rows = _read_rows(filename, raw)
    mapping = _map_columns(headers)
    rows, errors = _normalise_rows(raw_rows, mapping)
    if not rows:
        detail = errors[0]["reason"] if errors else "No transaction rows were found."
        raise ValueError(f"No usable transactions were found. First issue: {detail}")
    return analyse_rows(rows, filename, mapping, errors)
